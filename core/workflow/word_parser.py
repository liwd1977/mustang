"""AI-01：办事指南 Word 文档解析。"""

from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Iterator

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from PIL import Image

from core.workflow.flow_filter import (
    dedupe_sector_images,
    format_flow_display_title,
    is_flow_title_line,
    is_metadata_line,
    should_skip_flow_title,
    skip_reason_for,
)
from schemas.workflow import GuideParseResult, SectorBlock, WorkflowImage

CHINESE_NUM_MAP: dict[str, int] = {
    "一": 1,
    "二": 2,
    "三": 3,
    "四": 4,
    "五": 5,
    "六": 6,
    "七": 7,
    "八": 8,
    "九": 9,
    "十": 10,
}

SECTION_HEADING_RE = re.compile(
    r"^\s*([0-9]{1,2}|[一二三四五六七八九十]+)\s*[、．.\s]\s*(.+?)\s*$"
)

BLIP_XPATH = ".//*[local-name()='blip']"
REL_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
HEADING_STYLE_PREFIX = "Heading 1"
STACKED_IMAGE_MIN_HEIGHT = 1800


def parse_section_number(text: str) -> tuple[int, str] | None:
    text = text.strip()
    if not text:
        return None
    match = SECTION_HEADING_RE.match(text)
    if not match:
        return None
    raw_num, title = match.group(1), match.group(2).strip()
    if raw_num.isdigit():
        num = int(raw_num)
    elif raw_num in CHINESE_NUM_MAP:
        num = CHINESE_NUM_MAP[raw_num]
    else:
        return None
    return num, title


def _iter_block_items(parent: DocxDocument | Table | _Cell) -> Iterator[Paragraph | Table]:
    if isinstance(parent, Table):
        for row in parent.rows:
            for cell in row.cells:
                yield from _iter_block_items(cell)
        return

    if isinstance(parent, _Cell):
        for paragraph in parent.paragraphs:
            yield paragraph
        for table in parent.tables:
            yield from _iter_block_items(table)
        return

    body = parent.element.body
    for child in body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def _paragraph_text(paragraph: Paragraph) -> str:
    return paragraph.text.strip()


def _is_heading1(paragraph: Paragraph) -> bool:
    return paragraph.style.name.startswith(HEADING_STYLE_PREFIX) and bool(_paragraph_text(paragraph))


def _extract_images_from_paragraph(paragraph: Paragraph, doc: DocxDocument) -> list[bytes]:
    blobs: list[bytes] = []
    for blip in paragraph._element.xpath(BLIP_XPATH):
        rel_id = blip.get(REL_EMBED)
        if not rel_id or rel_id not in doc.part.related_parts:
            continue
        part = doc.part.related_parts[rel_id]
        if hasattr(part, "blob"):
            blobs.append(part.blob)
    return blobs


def _guess_extension(blob: bytes) -> str:
    if blob.startswith(b"\x89PNG"):
        return ".png"
    if blob.startswith(b"\xff\xd8"):
        return ".jpg"
    if blob.startswith(b"GIF"):
        return ".gif"
    return ".png"


def _split_image_blob(blob: bytes, parts: int) -> list[bytes]:
    """将纵向堆叠的多张流程图切分为独立图片。"""
    if parts <= 1:
        return [blob]
    with Image.open(io.BytesIO(blob)) as img:
        width, height = img.size
        part_h = max(1, height // parts)
        out: list[bytes] = []
        for i in range(parts):
            top = i * part_h
            bottom = height if i == parts - 1 else (i + 1) * part_h
            crop = img.crop((0, top, width, bottom))
            buf = io.BytesIO()
            crop.save(buf, format=img.format or "PNG")
            out.append(buf.getvalue())
        return out


def _image_height(blob: bytes) -> int:
    try:
        with Image.open(io.BytesIO(blob)) as img:
            return img.size[1]
    except Exception:
        return 0


def _save_image(blob: bytes, dest_dir: Path, index: int, *, title_hint: str = "") -> WorkflowImage:
    dest_dir.mkdir(parents=True, exist_ok=True)
    ext = _guess_extension(blob)
    filename = f"flow_{index:02d}{ext}"
    path = dest_dir / filename
    path.write_bytes(blob)
    width = height = 0
    try:
        with Image.open(io.BytesIO(blob)) as img:
            width, height = img.size
    except Exception:
        pass
    return WorkflowImage(
        index=index,
        filename=filename,
        path=str(path),
        width=width,
        height=height,
        title_hint=title_hint,
    )


def _append_flow_image(
    sector: SectorBlock,
    *,
    sector_dir: Path,
    doc: DocxDocument,
    blob: bytes | None,
    title_hint: str,
    pending_titles: list[str],
) -> None:
    if should_skip_flow_title(title_hint):
        sector.images.append(
            WorkflowImage(
                index=len(sector.images) + 1,
                title_hint=title_hint,
                skipped=True,
                skip_reason=skip_reason_for(title_hint),
            )
        )
        return
    if blob is None:
        return

    waiting = [t for t in pending_titles if t != title_hint]
    assign_titles = waiting + [title_hint]
    height = _image_height(blob)
    if height >= STACKED_IMAGE_MIN_HEIGHT and len(assign_titles) >= 2:
        for part_blob, hint in zip(_split_image_blob(blob, len(assign_titles)), assign_titles, strict=False):
            img = _save_image(
                part_blob,
                sector_dir / f"sector_{sector.index:02d}",
                len(sector.images) + 1,
                title_hint=hint,
            )
            sector.images.append(img)
        pending_titles.clear()
        return

    img = _save_image(
        blob,
        sector_dir / f"sector_{sector.index:02d}",
        len(sector.images) + 1,
        title_hint=title_hint,
    )
    sector.images.append(img)
    for title in assign_titles:
        while title in pending_titles:
            pending_titles.remove(title)


def parse_guide_docx(
    docx_path: Path,
    *,
    output_dir: Path,
    output_stem: str | None = None,
    section_from: int = 2,
    section_to: int = 10,
) -> GuideParseResult:
    docx_path = Path(docx_path)
    if not docx_path.is_file():
        return GuideParseResult(
            source_file=str(docx_path),
            errors=[f"文件不存在: {docx_path}"],
        )

    doc = Document(str(docx_path))
    folder_name = output_stem or docx_path.stem
    sector_dir = output_dir / folder_name
    sector_dir.mkdir(parents=True, exist_ok=True)

    sectors: dict[int, SectorBlock] = {}
    pending_titles_by_sector: dict[int, list[str]] = {}
    current_index: int | None = None
    current_title = ""
    text_lines: list[str] = []
    heading_counter = 0
    last_text_line = ""
    current_flow_title = ""
    current_flow_skipped = False
    line_before_metadata = ""

    def flush_sector() -> None:
        nonlocal current_index, current_title, text_lines, last_text_line
        nonlocal current_flow_title, current_flow_skipped, line_before_metadata
        if current_index is None:
            return
        if section_from <= current_index <= section_to:
            sectors[current_index] = SectorBlock(
                index=current_index,
                title=current_title,
                text_content="\n".join(text_lines).strip(),
                images=sectors.get(current_index, SectorBlock(index=current_index)).images,
            )
        text_lines = []
        last_text_line = ""
        current_flow_title = ""
        current_flow_skipped = False
        line_before_metadata = ""

    def _on_text_line(text: str) -> None:
        nonlocal last_text_line, current_flow_title, current_flow_skipped, line_before_metadata
        t = text.strip()
        last_text_line = t

        if is_metadata_line(t):
            if line_before_metadata and not is_flow_title_line(current_flow_title):
                current_flow_title = line_before_metadata
                current_flow_skipped = should_skip_flow_title(current_flow_title)
            return

        if is_flow_title_line(t):
            current_flow_title = t
            current_flow_skipped = should_skip_flow_title(t)
            line_before_metadata = ""
            if current_index is not None and not current_flow_skipped:
                pending_titles_by_sector.setdefault(current_index, []).append(t)
            return

        line_before_metadata = t

    def _image_title_hint() -> str:
        return current_flow_title or line_before_metadata or ""

    def _image_skipped() -> bool:
        return current_flow_skipped or should_skip_flow_title(_image_title_hint())

    for block in _iter_block_items(doc):
        if isinstance(block, Table):
            if current_index is not None and section_from <= current_index <= section_to:
                for row in block.rows:
                    for cell in row.cells:
                        for para in cell.paragraphs:
                            cell_text = _paragraph_text(para)
                            if cell_text:
                                text_lines.append(cell_text)
                                _on_text_line(cell_text)
                            for blob in _extract_images_from_paragraph(para, doc):
                                sector = sectors.setdefault(
                                    current_index,
                                    SectorBlock(index=current_index, title=current_title),
                                )
                                hint = _image_title_hint()
                                if _image_skipped():
                                    sector.images.append(
                                        WorkflowImage(
                                            index=len(sector.images) + 1,
                                            title_hint=hint,
                                            skipped=True,
                                            skip_reason=skip_reason_for(hint),
                                        )
                                    )
                                else:
                                    _append_flow_image(
                                        sector,
                                        sector_dir=sector_dir,
                                        doc=doc,
                                        blob=blob,
                                        title_hint=hint,
                                        pending_titles=pending_titles_by_sector.setdefault(current_index, []),
                                    )
            continue

        text = _paragraph_text(block)
        if text:
            if _is_heading1(block):
                flush_sector()
                heading_counter += 1
                current_index = heading_counter
                current_title = text
                if section_from <= current_index <= section_to:
                    sectors.setdefault(
                        current_index,
                        SectorBlock(index=current_index, title=current_title),
                    )
                continue

            heading = parse_section_number(text)
            if heading:
                flush_sector()
                current_index, current_title = heading
                if section_from <= current_index <= section_to:
                    sectors.setdefault(
                        current_index,
                        SectorBlock(index=current_index, title=current_title),
                    )
                continue

            if current_index is not None and section_from <= current_index <= section_to:
                text_lines.append(text)
                _on_text_line(text)

        if current_index is not None and section_from <= current_index <= section_to:
            for blob in _extract_images_from_paragraph(block, doc):
                sector = sectors.setdefault(
                    current_index,
                    SectorBlock(index=current_index, title=current_title),
                )
                hint = _image_title_hint()
                if _image_skipped():
                    sector.images.append(
                        WorkflowImage(
                            index=len(sector.images) + 1,
                            title_hint=hint,
                            skipped=True,
                            skip_reason=skip_reason_for(hint),
                        )
                    )
                else:
                    _append_flow_image(
                        sector,
                        sector_dir=sector_dir,
                        doc=doc,
                        blob=blob,
                        title_hint=hint,
                        pending_titles=pending_titles_by_sector.setdefault(current_index, []),
                    )

    flush_sector()

    ordered = [sectors[i] for i in sorted(sectors) if section_from <= i <= section_to]
    for sector in ordered:
        dedupe_sector_images(sector)
    return GuideParseResult(
        source_file=str(docx_path),
        output_dir=str(sector_dir),
        sectors=ordered,
    )


def parse_guide_docx_from_bytes(
    content: bytes,
    *,
    filename: str,
    output_dir: Path,
    section_from: int = 2,
    section_to: int = 10,
) -> GuideParseResult:
    temp_path = output_dir / "_upload" / filename
    temp_path.parent.mkdir(parents=True, exist_ok=True)
    temp_path.write_bytes(content)
    return parse_guide_docx(
        temp_path,
        output_dir=output_dir,
        section_from=section_from,
        section_to=section_to,
    )
